import csv
from io import BytesIO
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADERS = ["Business Name", "Phone", "Email", "Address", "Website", "Category", "Area", "City", "Province", "Country", "Latitude", "Longitude", "Map URL", "Place ID"]

def row_for(b):
    value = lambda key, default="": b.get(key, default) if isinstance(b, dict) else getattr(b, key, default)
    latitude, longitude = value("latitude", None), value("longitude", None)
    google_url = f"https://www.google.com/maps/search/?api=1&query={latitude},{longitude}" if latitude is not None and longitude is not None else value("google_maps_url")
    return [value("name"), value("phone") or "N/A", value("email") or "N/A", value("address") or "N/A", value("website") or "N/A", value("category") or "N/A", value("area") or "N/A", value("city"), value("province") or "N/A", value("country"),
            "" if latitude is None else latitude, "" if longitude is None else longitude, google_url or "N/A", value("place_id")]

def csv_response(businesses, filename):
    response = HttpResponse(content_type="text/csv; charset=utf-8")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response.write("\ufeff")
    writer = csv.writer(response)
    writer.writerow(HEADERS)
    for business in businesses:
        writer.writerow(row_for(business))
    return response

def excel_response(businesses, filename):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Businesses"
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
    for business in businesses:
        sheet.append(row_for(business))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [28, 20, 30, 45, 32, 22, 18, 18, 18, 18, 14, 14, 35, 28]
    for i, width in enumerate(widths, 1):
        from openpyxl.utils import get_column_letter
        sheet.column_dimensions[get_column_letter(i)].width = width
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
